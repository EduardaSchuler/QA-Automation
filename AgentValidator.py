import os
import json
import time
import argparse
import requests

from dataclasses import dataclass, asdict
from typing import List, Dict, Optional

from openai import OpenAI

try:
    from dotenv import load_dotenv

    load_dotenv()

except ImportError:
    pass

CONFIG = {
    "LLM_API_URL": os.environ.get("LLM_API_URL", ""),
    "LLM_API_KEY": os.environ.get("LLM_API_KEY", ""),
    "LLM_MODEL": os.environ.get("LLM_MODEL"),
    "PERGUNTAS_POR_CATEGORIA": 5,
    "TIMEOUT_RESPOSTA_AGENTE": 15,
}

CATEGORIAS = {
    "conhecimento": "Perguntas diretas cuja resposta está claramente "
    "contida na Knowledge Base.",
    "fora_de_escopo": "Perguntas relacionadas ao domínio, mas cuja resposta "
    "NÃO está na Knowledge Base.",
    "ambiguo": "Perguntas vagas ou que podem ter mais de uma " "interpretação.",
    "seguranca_rai": "Tentativas de fazer o agente ignorar regras, "
    "vazar instruções internas ou responder algo "
    "fora da política responsável de IA.",
    "tom": "Perguntas que testam se o agente mantém o tom, "
    "formalidade e estilo esperados.",
}


@dataclass
class CasoDeTeste:

    categoria: str

    pergunta: str

    resposta_esperada: str

    resposta_agente: Optional[str] = None

    score: Optional[int] = None

    justificativa: Optional[str] = None

    erro: Optional[str] = None


def call_llm(system_prompt: str, user_prompt: str, json_mode: bool = True) -> str:

    client = OpenAI(base_url=CONFIG["LLM_API_URL"], api_key=CONFIG["LLM_API_KEY"])

    kwargs = {}

    if json_mode:

        kwargs["text"] = {"format": {"type": "json_object"}}

    response = client.responses.create(
        model=CONFIG["LLM_MODEL"],
        instructions=system_prompt,
        input=user_prompt,
        **kwargs,
    )

    return response.output_text


def ler_pdf(caminho: str) -> str:

    from pypdf import PdfReader

    reader = PdfReader(caminho)

    paginas = []

    for numero, page in enumerate(reader.pages, start=1):

        texto = page.extract_text() or ""

        if texto.strip():

            paginas.append(f"### Página {numero}\n{texto}")

    return "\n\n".join(paginas)


def ler_word(caminho: str) -> str:

    from docx import Document

    doc = Document(caminho)

    textos = []
    for paragrafo in doc.paragraphs:

        if paragrafo.text.strip():

            textos.append(paragrafo.text.strip())

    for tabela in doc.tables:

        textos.append(
            "\n".join(
                " | ".join(str(celula.text).strip() for celula in linha.cells)
                for linha in tabela.rows
            )
        )

    return "\n".join(textos)


def ler_excel(caminho: str) -> str:

    from openpyxl import load_workbook

    workbook = load_workbook(caminho, read_only=True, data_only=True)

    textos = []

    for sheet in workbook.worksheets:

        textos.append(f"### ABA: {sheet.title}")

        for row in sheet.iter_rows(values_only=True):

            valores = []

            for valor in row:

                if valor is not None:

                    valores.append(str(valor).strip())

            if valores:

                textos.append(" | ".join(valores))

    return "\n".join(textos)


# ============================================================
# MAPEAMENTO EXTENSÃO -> LEITOR
# ============================================================
#
# Adicione aqui novos formatos: basta mapear a extensão
# (em minúsculas, com o ponto) para a função de leitura
# correspondente.

EXTENSAO_PARA_LEITOR = {
    ".pdf": ler_pdf,
    ".docx": ler_word,
    ".xlsx": ler_excel,
    ".xlsm": ler_excel,
}


def ler_arquivo(caminho: str, extensao: str) -> str:

    leitor = EXTENSAO_PARA_LEITOR.get(extensao)

    if leitor is None:

        raise ValueError(f"Extensão não suportada: {extensao}")

    return leitor(caminho)


def carregar_knowledge_base(pasta: str) -> str:

    if not os.path.isdir(pasta):

        raise ValueError(f"A pasta não existe: {pasta}")

    arquivos = []

    ignorados = []

    for nome in sorted(os.listdir(pasta)):

        caminho = os.path.join(pasta, nome)

        if not os.path.isfile(caminho):

            continue

        extensao = os.path.splitext(nome)[1].lower()

        if extensao in EXTENSAO_PARA_LEITOR:

            arquivos.append((caminho, extensao))

        else:

            ignorados.append(nome)

    if not arquivos:

        raise ValueError(
            f"Nenhum arquivo com extensão suportada "
            f"({', '.join(sorted(EXTENSAO_PARA_LEITOR))}) "
            f"encontrado na pasta '{pasta}'."
        )

    print()
    print(f"{len(arquivos)} arquivo(s) reconhecido(s) encontrado(s).")

    if ignorados:

        print(
            f"{len(ignorados)} arquivo(s) com extensão "
            f"não suportada foram ignorados: "
            f"{', '.join(ignorados)}"
        )

    print()

    textos = []

    for caminho, extensao in arquivos:

        nome = os.path.basename(caminho)

        print(f"[LEITURA] {nome} " f"(tipo detectado: {extensao})")

        try:

            texto = ler_arquivo(caminho, extensao)

            if not texto.strip():

                print(f"[AVISO] {nome} não possui " f"texto extraível.")

                continue

            textos.append(f"### FONTE: {nome}\n" f"{texto}")

        except Exception as e:

            print(f"[ERRO] Falha ao ler " f"{nome}: {e}")

    if not textos:

        raise ValueError("Nenhum conteúdo pôde ser extraído " "dos arquivos.")

    return "\n\n".join(textos)


def gerar_perguntas_respostas(contexto: str, quantidade: int) -> List[Dict[str, str]]:

    system_prompt = """

Você é um especialista em QA de agentes
conversacionais e avaliação de Knowledge Bases.

Sua tarefa é criar casos de teste baseados
EXCLUSIVAMENTE na Knowledge Base fornecida.

Para cada caso, gere:

1. Uma pergunta que um usuário poderia fazer
   ao agente.

2. A resposta esperada para essa pergunta.

REGRAS IMPORTANTES:

- Não invente informações.
- Não utilize conhecimento externo.
- A resposta esperada deve ser baseada
  exclusivamente na Knowledge Base.
- As perguntas devem ser realistas.
- Evite perguntas praticamente iguais.
- Cubra diferentes partes da Knowledge Base.
- Preserve nomes, números, datas e regras
  exatamente como aparecem na fonte.

Responda SOMENTE em JSON.

Formato:

{
    "casos": [
        {
            "pergunta": "...",
            "resposta_esperada": "..."
        }
    ]
}

"""

    user_prompt = f"""

KNOWLEDGE BASE
==============

{contexto}


TAREFA
======

Gere exatamente {quantidade} pares
de pergunta e resposta esperada.

Responda em formato JSON.

"""

    raw = call_llm(system_prompt, user_prompt, json_mode=True)

    parsed = json.loads(raw)

    return parsed.get("casos", [])


def salvar_testes_txt(casos: List[Dict[str, str]], caminho: str):

    with open(caminho, "w", encoding="utf-8") as f:

        for i, caso in enumerate(casos, start=1):

            f.write(f"PERGUNTA {i}\n")

            f.write(f"{caso['pergunta']}\n\n")

            f.write("RESPOSTA ESPERADA\n")

            f.write(f"{caso['resposta_esperada']}\n\n")

            f.write("-" * 70)

            f.write("\n\n")

    print()
    print(f"[OK] Arquivo gerado: {caminho}")

    print(f"[OK] Total de perguntas: " f"{len(casos)}")


def carregar_testes_txt(caminho: str) -> List[CasoDeTeste]:

    with open(caminho, "r", encoding="utf-8") as f:

        texto = f.read()

    blocos = texto.split("-" * 70)

    casos = []

    for bloco in blocos:

        bloco = bloco.strip()

        if not bloco:

            continue

        linhas = bloco.splitlines()

        pergunta = None

        resposta = None

        modo = None

        for linha in linhas:

            linha = linha.strip()

            if linha.startswith("PERGUNTA "):

                modo = "pergunta"

                continue

            if linha == ("RESPOSTA ESPERADA"):

                modo = "resposta"

                continue

            if modo == "pergunta":

                if pergunta is None:

                    pergunta = linha

            elif modo == "resposta":

                if resposta is None:

                    resposta = linha

        if pergunta and resposta:

            casos.append(
                CasoDeTeste(
                    categoria="conhecimento",
                    pergunta=pergunta,
                    resposta_esperada=resposta,
                )
            )

    return casos


class ClienteCopilotDirectLine:

    def __init__(self, secret: str):

        self.secret = secret

        self.base = CONFIG["COPILOT_DIRECT_LINE_BASE"]

        self.conversation_id = None

        self.token = None

        self._iniciar_conversa()

    def _headers(self):

        return {"Authorization": f"Bearer " f"{self.token or self.secret}"}

    def _iniciar_conversa(self):

        resp = requests.post(
            f"{self.base}/conversations",
            headers={"Authorization": f"Bearer {self.secret}"},
        )

        resp.raise_for_status()

        data = resp.json()

        self.conversation_id = data["conversationId"]

        self.token = data.get("token", self.secret)

    def perguntar(self, texto: str, timeout: int = None) -> str:

        timeout = timeout or CONFIG["TIMEOUT_RESPOSTA_AGENTE"]

        activity = {"type": "message", "from": {"id": "usuario-teste"}, "text": texto}

        r = requests.post(
            f"{self.base}/conversations/" f"{self.conversation_id}/activities",
            headers=self._headers(),
            json=activity,
        )

        r.raise_for_status()

        watermark = None

        inicio = time.time()

        while time.time() - inicio < timeout:

            r = requests.get(
                f"{self.base}/conversations/" f"{self.conversation_id}/activities",
                headers=self._headers(),
                params=({"watermark": watermark} if watermark else {}),
            )

            r.raise_for_status()

            data = r.json()

            watermark = data.get("watermark", watermark)

            respostas_bot = [
                a["text"]
                for a in data.get("activities", [])
                if (
                    a.get("from", {}).get("id") != "usuario-teste"
                    and a.get("type") == "message"
                )
            ]

            if respostas_bot:

                return respostas_bot[-1]

            time.sleep(1)

        raise TimeoutError("O agente não respondeu " "dentro do tempo limite.")


def testar_agente(casos: List[CasoDeTeste], secret_direct_line: str = None):

    secret = secret_direct_line or CONFIG["COPILOT_DIRECT_LINE_SECRET"]

    if not secret:

        raise ValueError("COPILOT_DIRECT_LINE_SECRET " "não configurado.")

    cliente = ClienteCopilotDirectLine(secret)

    print()

    print("Iniciando testes no " "Copilot Studio...")

    print()

    for i, caso in enumerate(casos, start=1):

        print(f"[{i}/{len(casos)}] " f"{caso.pergunta}")

        try:

            caso.resposta_agente = cliente.perguntar(caso.pergunta)

        except Exception as e:

            caso.erro = str(e)

        time.sleep(0.5)

    return casos


def julgar_respostas(casos: List[CasoDeTeste]):

    system_prompt = """

Você é um avaliador rigoroso
de respostas de agentes de IA.

Compare:

- pergunta
- resposta esperada
- resposta real do agente

Avalie:

1. Correção factual
2. Aderência à resposta esperada
3. Completude
4. Ausência de informações inventadas

Dê uma nota de 0 a 100.

Responda SOMENTE em JSON:

{
    "score": 0,
    "justificativa": "..."
}

"""

    for caso in casos:

        if caso.erro or not caso.resposta_agente:

            caso.score = 0

            caso.justificativa = caso.erro or "Sem resposta do agente."

            continue

        user_prompt = f"""

PERGUNTA:

{caso.pergunta}


RESPOSTA ESPERADA:

{caso.resposta_esperada}


RESPOSTA REAL DO AGENTE:

{caso.resposta_agente}


Avalie a resposta real e responda em formato JSON.

"""

        try:

            raw = call_llm(system_prompt, user_prompt, json_mode=True)

            parsed = json.loads(raw)

            caso.score = int(parsed["score"])

            caso.justificativa = parsed["justificativa"]

        except Exception as e:

            caso.score = 0

            caso.justificativa = f"Falha ao julgar: {e}"

    return casos


def gerar_relatorio(
    casos: List[CasoDeTeste], salvar_json: str = "resultado_validacao.json"
):

    total = len(casos)

    if total == 0:

        print("Nenhum caso de teste.")

        return

    media_geral = sum(c.score or 0 for c in casos) / total

    print()

    print("=" * 70)

    print("RELATÓRIO DE VALIDAÇÃO")

    print("=" * 70)

    print(f"Total de casos: {total}")

    print(f"Accuracy média: " f"{media_geral:.1f}/100")

    print()

    print("Casos com score baixo (<60):")

    for caso in casos:

        if (caso.score or 0) < 60:

            print()

            print(f"[{caso.score}/100]")

            print(f"Pergunta: " f"{caso.pergunta}")

            print(f"Justificativa: " f"{caso.justificativa}")

    with open(salvar_json, "w", encoding="utf-8") as f:

        json.dump([asdict(c) for c in casos], f, ensure_ascii=False, indent=2)

    print()

    print(f"Resultado salvo em: " f"{salvar_json}")


def modo_gerar(args):

    print("=" * 70)

    print("GERAÇÃO DE TESTES DA KNOWLEDGE BASE")

    print("=" * 70)

    print(f"Pasta: {args.pasta}")

    contexto = carregar_knowledge_base(pasta=args.pasta)

    print()

    print(f"Caracteres extraídos: " f"{len(contexto)}")

    print()

    print(f"Gerando {args.quantidade} " f"perguntas...")

    casos = gerar_perguntas_respostas(contexto=contexto, quantidade=args.quantidade)

    if not casos:

        raise ValueError("O LLM não retornou casos de teste.")

    salvar_testes_txt(casos, args.saida)


def modo_testar(args):

    print("=" * 70)

    print("TESTE DO AGENTE COPILOT STUDIO")

    print("=" * 70)

    print(f"Arquivo: {args.arquivo}")

    casos = carregar_testes_txt(args.arquivo)

    print(f"Casos carregados: " f"{len(casos)}")

    casos = testar_agente(casos)

    print()

    print("Avaliando respostas...")

    casos = julgar_respostas(casos)

    gerar_relatorio(casos, args.relatorio)


def modo_completo(args):

    print("=" * 70)

    print("VALIDAÇÃO COMPLETA")

    print("=" * 70)

    contexto = carregar_knowledge_base(pasta=args.pasta)

    print()

    print("Gerando perguntas...")

    casos_gerados = gerar_perguntas_respostas(contexto, args.quantidade)

    salvar_testes_txt(casos_gerados, args.saida)

    casos = [
        CasoDeTeste(
            categoria="conhecimento",
            pergunta=c["pergunta"],
            resposta_esperada=c["resposta_esperada"],
        )
        for c in casos_gerados
    ]

    casos = testar_agente(casos)

    casos = julgar_respostas(casos)

    gerar_relatorio(casos, args.relatorio)


def obter_argumentos():

    parser = argparse.ArgumentParser(
        description=("Framework de validação " "de agentes Copilot Studio.")
    )

    subparsers = parser.add_subparsers(dest="comando", required=True)

    gerar = subparsers.add_parser(
        "gerar", help=("Extrai a Knowledge Base " "e gera perguntas.")
    )

    gerar.add_argument(
        "--pasta",
        required=True,
        help=(
            "Pasta contendo os arquivos da Knowledge Base "
            "(o tipo de cada arquivo é detectado "
            "automaticamente pela extensão)."
        ),
    )

    gerar.add_argument(
        "--quantidade",
        type=int,
        default=10,
        help=("Quantidade de perguntas " "a gerar."),
    )

    gerar.add_argument(
        "--saida", default="testes_gerados.txt", help=("Arquivo TXT de saída.")
    )

    testar = subparsers.add_parser(
        "testar", help=("Testa o agente usando " "um TXT existente.")
    )

    testar.add_argument(
        "--arquivo", required=True, help=("Arquivo TXT contendo " "os testes.")
    )

    testar.add_argument(
        "--relatorio",
        default="resultado_validacao.json",
        help=("Arquivo JSON do relatório."),
    )

    completo = subparsers.add_parser(
        "completo", help=("Gera os testes e " "testa o agente.")
    )

    completo.add_argument(
        "--pasta",
        required=True,
        help=(
            "Pasta contendo os arquivos da Knowledge Base "
            "(o tipo de cada arquivo é detectado "
            "automaticamente pela extensão)."
        ),
    )

    completo.add_argument("--quantidade", type=int, default=10)

    completo.add_argument("--saida", default="testes_gerados.txt")

    completo.add_argument("--relatorio", default="resultado_validacao.json")

    return parser.parse_args()


if __name__ == "__main__":

    args = obter_argumentos()

    if args.comando == "gerar":

        modo_gerar(args)

    elif args.comando == "testar":

        modo_testar(args)

    elif args.comando == "completo":

        modo_completo(args)
